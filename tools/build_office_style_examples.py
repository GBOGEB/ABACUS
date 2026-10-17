"""Generate privacy-safe real DOCX/XLSX/PPTX Office style inputs.

Synthetic engineering-style illustrations only; not ABACUS CI receipts or
QPS commercial values. Requires python-docx and openpyxl; PPTX is generated
by sibling PptxGenJS script (node runtime with pptxgenjs available).
No external templates or sensitive uploads are embedded.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import subprocess

NAVY = "17283F"
BLUE = "216BB0"
TEAL = "138D88"
AMBER = "D88731"
LIGHT = "ECF4F7"


def create_docx(path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.80)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 25, NAVY), ("Heading 1", 15, BLUE), ("Heading 2", 11.5, TEAL)]:
        st = doc.styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(13)
        st.paragraph_format.space_after = Pt(7)
    kicker = doc.styles.add_style("Control Kicker", WD_STYLE_TYPE.PARAGRAPH)
    kicker.font.name = "Aptos"
    kicker.font.size = Pt(9)
    kicker.font.bold = True
    kicker.font.color.rgb = RGBColor.from_string(TEAL)
    doc.core_properties.title = "CODE PURPLE | Office Style Input"
    doc.core_properties.subject = "Synthetic style and format extraction fixture"
    doc.core_properties.author = "CODE PURPLE synthetic fixtures"
    header = section.header.paragraphs[0]
    header.text = "CODE PURPLE    /    OFFICE REFERENCE SERIES"
    header.style = kicker
    doc.add_paragraph("CONTROLLED TEST INPUT  /  SYNTHETIC DATA", style="Control Kicker")
    doc.add_heading("Format, style & evidence", 0)
    p = doc.add_paragraph("A realistic Word document for verifying ")
    run = p.add_run("heading inheritance, emphasis, tables and traceable structure")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(" without exposing project evidence or commercial source content.")
    doc.add_heading("1  Operational scope", 1)
    p = doc.add_paragraph("Input quality is checked at the producer boundary. Derived style evidence shall not replace authoritative source bytes or permitted review controls.")
    doc.add_heading("1.1  Expected style features", 2)
    for x in ["Named Heading 1 / Heading 2 styles with a theme font", "An actual Word table with a colored header and mixed alignment", "Footer, page margins, explicit bold/italic/color and an evidence note"]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("2  Example review matrix", 1)
    headers = ["ID", "Input family", "Control", "Example state"]
    rows = [["T01", "DOCX", "Heading + table", "OBSERVED"],
            ["T02", "XLSX", "Formula + number format", "OBSERVED"],
            ["T03", "PPTX", "Theme + shape", "OBSERVED"]]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    table.autofit = False
    widths = [Inches(0.62), Inches(1.07), Inches(2.75), Inches(1.25)]
    for i,h in enumerate(headers):
        table.rows[0].cells[i].text = h
        table.rows[0].cells[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i,txt in enumerate(row):
            cells[i].text = txt
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shading)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
    doc.add_heading("3  Evidence & privacy", 1)
    doc.add_paragraph("All counts and status labels above are demonstration content. Style extraction never publishes body text, sheet names, speaker notes, source paths, or a binary Office file to a public repository.")
    p = doc.add_paragraph()
    r = p.add_run("Illustrative evidence only — not an accepted QPS record.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(AMBER)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "Synthetic fixture  |  CODE PURPLE / v1"
    doc.save(path)


def create_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties
    wb = Workbook()
    cover = wb.active
    cover.title = "Control Board"
    runs = wb.create_sheet("Test Runs")
    guide = wb.create_sheet("Style Guide")
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    for s in (cover, runs, guide):
        s.sheet_view.showGridLines = False
        s.freeze_panes = "B5" if s is runs else "B4"
        s.sheet_properties.pageSetUpPr.fitToPage = True
    cover.merge_cells("B2:H3")
    cover["B2"] = "CODE PURPLE  /  TEST EVIDENCE"
    cover["B2"].font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    cover["B2"].fill = PatternFill("solid", fgColor=NAVY)
    cover["B2"].alignment = Alignment(vertical="center")
    cover.row_dimensions[2].height = 26
    cover.row_dimensions[3].height = 26
    for col in "ABCDEFGH":
        cover.column_dimensions[col].width = 16
    cover.column_dimensions["B"].width = 28
    cover["B5"] = "Synthetic fixture only — not real pipeline evidence"
    cover["B5"].font = Font(name="Aptos", color=AMBER, italic=True, size=10)
    for cell, label in [("B7", "Observed runs"), ("B8", "Total passed"), ("B9", "Total failed"), ("B10", "Mean duration (s)")]:
        cover[cell] = label
        cover[cell].font = Font(name="Aptos", size=11, color=NAVY, bold=True)
    for cell, formula in [("E7", "=COUNTA('Test Runs'!B5:B10)"),
                          ("E8", "=SUM('Test Runs'!D5:D10)"),
                          ("E9", "=SUM('Test Runs'!E5:E10)"),
                          ("E10", "=AVERAGE('Test Runs'!G5:G10)")]:
        cover[cell] = formula
        cover[cell].font = Font(name="Aptos", size=12, bold=True, color="000000")
        cover[cell].fill = PatternFill("solid", fgColor=LIGHT)
        cover[cell].number_format = "0.0" if cell == "E10" else "0"
    runs.merge_cells("A1:G2")
    runs["A1"] = "EXAMPLE TEST HISTORY  |  SIX REVISION-BOUND OBSERVATIONS"
    runs["A1"].font = Font(name="Aptos Display", size=17, bold=True, color="FFFFFF")
    runs["A1"].fill = PatternFill("solid", fgColor=NAVY)
    runs["A1"].alignment = Alignment(vertical="center")
    header = ["Wave", "Revision", "Status", "Passed", "Failed", "Skipped", "Seconds"]
    for i,h in enumerate(header,1):
        cell = runs.cell(4,i,h)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center")
    samples = [
        ("N1", "rev-a", "PASS", 52, 0, 2, 6.8),
        ("N2", "rev-b", "PASS", 53, 0, 2, 6.4),
        ("N3", "rev-c", "FAIL", 53, 1, 2, 5.2),
        ("N4", "rev-d", "FAIL", 54, 2, 2, 7.0),
        ("N5", "rev-e", "FAIL", 56, 1, 2, 4.9),
        ("N6", "rev-f", "PASS", 58, 0, 2, 4.6),
    ]
    for row_index, row in enumerate(samples,5):
        for ci,value in enumerate(row,1):
            cell = runs.cell(row_index,ci,value)
            cell.font = Font(name="Aptos", size=10.5, color=("0000FF" if ci in (4,5,6,7) else "606B78"))
            if row_index%2==0:
                cell.fill = PatternFill("solid", fgColor="F3F8FA")
            if ci>=4:
                cell.number_format = "0.0" if ci==7 else "0"
    runs.column_dimensions["A"].width = 12
    runs.column_dimensions["B"].width = 18
    runs.column_dimensions["C"].width = 16
    for col in "DEFG":
        runs.column_dimensions[col].width = 16
    runs.conditional_formatting.add("E5:E10", CellIsRule(operator="greaterThan",formula=["0"],fill=PatternFill("solid",fgColor="FCE9E5")))
    t = Table(displayName="SyntheticRuns",ref="A4:G10")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",showRowStripes=False,showColumnStripes=False)
    runs.add_table(t)
    ch = BarChart()
    ch.title = "Passed tests per revision"
    ch.y_axis.title = "Tests"
    ch.x_axis.title = "Revision"
    ch.style = 10
    ch.height, ch.width = 7, 15
    ch.add_data(Reference(runs,min_col=4,min_row=4,max_row=10),titles_from_data=True)
    ch.set_categories(Reference(runs,min_col=2,min_row=5,max_row=10))
    runs.add_chart(ch,"I4")
    guide["B2"] = "STYLE SYSTEM / INPUT CONTRACT"
    guide["B2"].font = Font(name="Aptos Display",size=19,bold=True,color=NAVY)
    guide["B4"] = "Token"
    guide["C4"] = "Value"
    guide["D4"] = "Usage"
    for cell in guide[4][1:4]:
        cell.fill=PatternFill("solid",fgColor=NAVY)
        cell.font=Font(color="FFFFFF",bold=True)
    for i,row in enumerate([
        ("navy", "#"+NAVY, "titles/control"),
        ("blue", "#"+BLUE, "table heads"),
        ("teal", "#"+TEAL, "success/structure"),
        ("amber", "#"+AMBER, "caution/review"),
    ],5):
        for j,val in enumerate(row,2):
            guide.cell(i,j,val)
            guide.cell(i,j).font=Font(name="Aptos",size=11,color="606B78")
        guide.cell(i,3).fill=PatternFill("solid",fgColor=row[1].lstrip("#"))
        guide.cell(i,3).font=Font(color="FFFFFF",bold=True)
    guide.column_dimensions["B"].width=24
    guide.column_dimensions["C"].width=19
    guide.column_dimensions["D"].width=31
    wb.properties.title="CODE PURPLE Office Style Synthetic Workbook"
    wb.properties.creator="CODE PURPLE synthetic fixtures"
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--no-deck", action="store_true")
    args = parser.parse_args()
    out = args.out
    out.mkdir(parents=True, exist_ok=True)
    files = [out/"CODE_PURPLE_Office_Style_Sample.docx", out/"CODE_PURPLE_Office_Style_Sample.xlsx"]
    create_docx(files[0]); create_xlsx(files[1])
    if not args.no_deck:
        deck = out/"CODE_PURPLE_Office_Style_Sample.pptx"
        subprocess.run(["node", str(Path(__file__).with_name("build_office_style_deck.js")), str(deck)],check=True)
        files.append(deck)
    manifest={"schema":"abacus.synthetic-office-inputs.v1","authority_transfer":False,
              "commercial_qps_data":False,"purpose":"style_and_format_input_reference",
              "source_class":"SYNTHETIC_DEMONSTRATION", "files":[]}
    for file in files:
        manifest["files"].append({"name":file.name,"bytes":file.stat().st_size,
                                  "sha256":hashlib.sha256(file.read_bytes()).hexdigest()})
    (out/"office_input_manifest.json").write_text(json.dumps(manifest,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps(manifest,indent=2))

if __name__ == "__main__":
    main()